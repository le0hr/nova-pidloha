from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.config import Settings
from app.parser.values import identifier
from app.parser.workbook import parse_workbook


def workbook(tmp_path: Path, attrs: list[str] | None = None, price: object = 10.25) -> Path:
    attrs = attrs or []
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Двері"
    ws.append(["ID товару", "Назва"])
    ws.append([1020, "Door"])
    ws.append([])
    ws.append(["product_id", "ID Варіації", *attrs, "Ціна", "SKU"])
    ws.append([1020, 102001, *[f"value{i}" for i in range(len(attrs))], price, "SKU-1"])
    path = tmp_path / "catalog.xlsx"
    wb.save(path)
    return path


@pytest.mark.parametrize(
    "attrs",
    [
        [],
        ["Колір"],
        ["Декор", "Колір", "Скло", "Кромка", "Тип замка", "Сторона відкривання"],
        [f"New {i}" for i in range(15)],
    ],
)
def test_dynamic_attributes(tmp_path: Path, attrs: list[str]) -> None:
    result = parse_workbook(workbook(tmp_path, attrs))
    assert result.valid, result.issues
    v = result.catalog.variations[0]
    assert len(v.attributes) == len(attrs)
    assert [a.display_name for a in v.attribute_definitions] == attrs
    assert v.price == Decimal("10.25")
    assert v.external_id == "102001"
    assert v.sku == "SKU-1"


@pytest.mark.parametrize("price", ["abc", -1, None, "NaN", True])
def test_invalid_price(tmp_path: Path, price: object) -> None:
    assert not parse_workbook(workbook(tmp_path, price=price)).valid


def edit(path: Path, cell: str, value: object) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    assert ws is not None
    assert value is None or isinstance(value, (str, int, float))
    ws[cell] = value
    wb.save(path)
    wb.close()


@pytest.mark.parametrize(
    "cell,value",
    [
        ("A5", 999),
        ("B5", None),
        ("B2", None),
        ("C4", "Prce"),
        ("C4", None),
        ("B4", "price"),
        ("C5", "=1+2"),
    ],
)
def test_bad_rows_and_headers(tmp_path: Path, cell: str, value: object) -> None:
    p = workbook(tmp_path)
    edit(p, cell, value)
    assert not parse_workbook(p).valid


@pytest.mark.parametrize("kind", ["product", "variation"])
def test_duplicates(tmp_path: Path, kind: str) -> None:
    p = workbook(tmp_path)
    from openpyxl import load_workbook

    wb = load_workbook(p)
    ws = wb.active
    assert ws is not None
    if kind == "product":
        ws.insert_rows(3)
        ws.cell(3, 1, 1020)
        ws.cell(3, 2, "Other")
    else:
        ws.append([1020, 102001, 30])
    wb.save(p)
    assert not parse_workbook(p).valid


def test_horizontal_sections_and_multiple_sheets(tmp_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.worksheets[0])
    for sheet, pid in [("Двері", 1020), ("Підлога", 2001)]:
        ws = wb.create_sheet(sheet)
        ws.append(
            [
                " Назва ",
                " PRODUCT   ID ",
                None,
                "ID продукту",
                "Характеристика",
                "Значення",
                "Одиниця",
                None,
                "product_id",
                "addon",
                "price",
            ]
        )
        ws.append(["Name", pid, None, pid, "Width", 80, "cm", None, pid, "Handle", "25,50"])
        ws.append([None, None, None, pid, "Height", 200, "cm"])
    p = tmp_path / "horizontal.xlsx"
    wb.save(p)
    result = parse_workbook(p)
    assert result.valid, result.issues
    assert len(result.catalog.products) == 2
    assert len(result.catalog.products[0].characteristics) == 2
    assert result.catalog.products[0].addons[0].price == Decimal("25.50")


def test_empty_attribute_header(tmp_path: Path) -> None:
    p = workbook(tmp_path, ["Color"])
    edit(p, "C4", None)
    assert not parse_workbook(p).valid


def test_reserved_attribute(tmp_path: Path) -> None:
    p = workbook(tmp_path, ["active"])
    assert not parse_workbook(p).valid


def test_attribute_mapping(tmp_path: Path) -> None:
    result = parse_workbook(
        workbook(tmp_path, ["Колір"]), Settings(attribute_keys={"Колір": "color"})
    )
    assert result.catalog.variations[0].attributes == {"color": "value0"}


def test_numeric_id() -> None:
    assert identifier(102001.0) == "102001"
    assert identifier("00102") == "00102"


def test_blank_rows(tmp_path: Path) -> None:
    p = workbook(tmp_path)
    edit(p, "A7", 1020)
    edit(p, "B7", 102002)
    edit(p, "C7", 5)
    result = parse_workbook(p)
    assert result.valid, result.issues
    assert any(i.severity == "WARNING" for i in result.issues)


def test_malformed_side_section(tmp_path: Path) -> None:
    p = workbook(tmp_path)
    edit(p, "G1", "product_id")
    edit(p, "H1", "characteristic")
    edit(p, "G2", 1020)
    edit(p, "H2", "Width")
    assert not parse_workbook(p).valid


def test_documented_workbook_integration() -> None:
    result = parse_workbook(Path(__file__).parents[1] / "examples" / "catalog.xlsx")
    assert result.valid, result.issues
    assert len(result.catalog.products) == 2
    assert len(result.catalog.variations) == 4
    assert len(result.catalog.filters) == 2
    assert result.catalog.products[0].description == "Опис товару"
    assert len(result.catalog.products[0].characteristics) == 2


@pytest.mark.parametrize("section", ["characteristic", "addon", "filter"])
def test_missing_side_reference(tmp_path: Path, section: str) -> None:
    p = workbook(tmp_path)
    for cell, value in [
        ("G1", "product_id"),
        ("H1", section),
        ("I1", "value" if section != "addon" else "price"),
        ("G2", "missing"),
        ("H2", "Example"),
        ("I2", 10),
    ]:
        edit(p, cell, value)
    assert not parse_workbook(p).valid


def test_duplicate_combination(tmp_path: Path) -> None:
    p = workbook(tmp_path, ["Колір"])
    for cell, value in [("A6", 1020), ("B6", 102002), ("C6", "value0"), ("D6", 20)]:
        edit(p, cell, value)
    result = parse_workbook(p)
    assert result.valid, result.issues
    assert [v.external_id for v in result.catalog.variations] == ["102001", "102002"]
    assert [v.price for v in result.catalog.variations] == [Decimal("10.25"), Decimal(20)]
    assert any(i.field == "attributes" and i.severity == "WARNING" for i in result.issues)


def test_cross_sheet_reference(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    p = workbook(tmp_path)
    w = load_workbook(p)
    s = w.create_sheet("Інша категорія")
    s.append(["product_id", "characteristic", "value"])
    s.append([1020, "Width", 80])
    w.save(p)
    result = parse_workbook(p)
    assert not result.valid
    assert any("same category" in i.message for i in result.issues)


def test_duplicate_headers(tmp_path: Path) -> None:
    p = workbook(tmp_path, ["Color", " color "])
    assert not parse_workbook(p).valid


def test_price_before_id(tmp_path: Path) -> None:
    p = workbook(tmp_path)
    edit(p, "B4", "price")
    edit(p, "C4", "variation_id")
    assert not parse_workbook(p).valid


def test_alias_and_sku_configuration(tmp_path: Path) -> None:
    p = workbook(tmp_path)
    edit(p, "B4", "SKU")
    edit(p, "D4", "note")
    edit(p, "C4", "Вартість")
    result = parse_workbook(
        p, Settings(sku_is_variation_id=True, header_aliases={"price": ["Вартість"]})
    )
    assert result.valid, result.issues


def test_unknown_after_price_rejected(tmp_path: Path) -> None:
    p = workbook(tmp_path)
    edit(p, "D4", "New Attribute")
    assert not parse_workbook(p).valid


def test_actual_personnel_workbook() -> None:
    result = parse_workbook(Path(__file__).parents[1] / "example.xlsx")
    assert not result.valid
    assert len(result.catalog.products) == 2
    assert len(result.catalog.variations) == 58
    assert len(result.catalog.products[0].characteristics) == 2
    assert result.catalog.products[0].addons[0].group == "Коробка"
    assert result.catalog.products[0].addons[0].price_delta == Decimal(1020)
    assert result.catalog.products[0].short_description is not None
    assert not any(i.field == "header" for i in result.issues)
    assert any(i.source.sheet == "Підлога" and i.field == "name" for i in result.issues)
    assert sum(i.field == "attributes" for i in result.issues) == 8
    assert any(i.field == "image" for i in result.issues)


def test_corrected_personnel_workbook(tmp_path: Path) -> None:
    from example_support import corrected_example

    result = parse_workbook(corrected_example(tmp_path))
    assert result.valid, result.issues
    assert len(result.catalog.products) == 3
    assert len(result.catalog.variations) == 50
    floor = next(p for p in result.catalog.products if p.external_id == "2001")
    assert floor.price == Decimal(100)
    assert floor.attributes["площа_в_пачці"] == 5.2
    assert len(result.assets) == 1


def test_repeated_combined_details_conflict(tmp_path: Path) -> None:
    from example_support import corrected_example

    path = corrected_example(tmp_path)
    from openpyxl import load_workbook

    w = load_workbook(path)
    w["Двері"]["D4"] = "Conflicting product name"
    w.save(path)
    assert any("Conflicting details" in i.message for i in parse_workbook(path).issues)


def test_missing_photo_blocks_import(tmp_path: Path) -> None:
    path = workbook(tmp_path)
    edit(path, "D4", "image")
    edit(path, "D5", "missing.jpg")
    assert any(i.field == "image" for i in parse_workbook(path).issues)


def test_cost_header_maps_to_price(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["ID товару", "Назва", "вартість"])
    ws.append([4001, "Handle", "1052,50"])
    path = tmp_path / "cost.xlsx"
    wb.save(path)
    result = parse_workbook(path)
    assert result.valid, result.issues
    assert result.catalog.products[0].price == Decimal("1052.50")
    assert "вартість" not in result.catalog.products[0].attributes


def test_wrong_filter_header_suggests_product_without_retargeting(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append([None, None, None, "Фільтри"])
    ws.append(["ID товару", "Назва", None, "ID варіації", "Назва", "Значення"])
    ws.append([1020, "Handle", None, 1020, "Color", "White"])
    path = tmp_path / "filters.xlsx"
    wb.save(path)
    result = parse_workbook(path)
    assert not result.valid
    assert any('Change "ID варіації" to "ID товару"' in i.message for i in result.issues)
    assert result.catalog.filters[0].target_type == "variation"
    edit(path, "D2", "ID товару")
    assert parse_workbook(path).valid


def test_blank_variation_warnings_are_grouped(tmp_path: Path) -> None:
    path = workbook(tmp_path)
    edit(path, "A100", 1020)
    edit(path, "B100", 102002)
    edit(path, "C100", 5)
    result = parse_workbook(path)
    warnings = [i for i in result.issues if i.severity == "WARNING"]
    assert result.valid, result.issues
    assert len(warnings) == 1
    assert "94 empty variation rows" in warnings[0].message


def test_duplicate_id_reports_original_sheet(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    path = workbook(tmp_path)
    wb = load_workbook(path)
    ws = wb.create_sheet("Hardware")
    ws.append(["product_id", "name"])
    ws.append([1020, "Handle"])
    wb.save(path)
    result = parse_workbook(path)
    assert not result.valid
    assert any('First defined on sheet "Двері", row 2' in i.message for i in result.issues)


@pytest.mark.parametrize("photo", [r"C:\Users\person\photo.jpg", r"\\server\photos\photo.jpg"])
def test_windows_photo_path_has_actionable_error(tmp_path: Path, photo: str) -> None:
    path = workbook(tmp_path)
    edit(path, "D4", "image")
    edit(path, "D5", photo)
    result = parse_workbook(path)
    assert not result.valid
    assert any("Windows drive and network paths" in i.message for i in result.issues)


def test_symbol_filter_values_have_distinct_stable_urls() -> None:
    from app.strapi.mapping import filter_slug

    symbols = ["●", "▲", "▪"]
    slugs = [filter_slug(symbol, "Позначка каталогу") for symbol in symbols]
    assert len(set(slugs)) == 3
    assert all(slug.startswith("позначка-каталогу-symbol-") for slug in slugs)
    assert slugs == [filter_slug(symbol, "Позначка каталогу") for symbol in symbols]
    assert filter_slug("Дуб білений", "Колір") == "колір-дуб-білений"


def test_independent_variation_addons_and_filters(tmp_path: Path) -> None:
    from app.strapi.mapping import desired_records

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Doors"
    ws.append(
        [None, None, None, None, None, "Додатки", None, None, None, None, "Фільтри", None, None]
    )
    ws.merge_cells("F1:I1")
    ws.merge_cells("K1:M1")
    ws.append(
        [
            "ID товару",
            "Назва",
            "ID варіації",
            "Ціна",
            None,
            "ID варіації",
            "Клас",
            "Назва",
            "Збільшення ціни",
            None,
            "ID варіації",
            "Назва",
            "Значення",
        ]
    )
    ws.append(["p1", "Door", "v1", 100, None, "v1", "Frame", "Oak", 10, None, "v1", "Color", "Oak"])
    ws.append(["p1", "Door", "v2", 110, None, "v1", "Frame", "White", 20])
    ws.append([None, None, None, None, None, "v2", "Frame", "Oak", 30])
    path = tmp_path / "independent.xlsx"
    wb.save(path)
    result = parse_workbook(path)
    assert result.valid, result.issues
    assert len(result.tables) == 3
    assert not result.catalog.products[0].addons
    assert [len(v.addons) for v in result.catalog.variations] == [2, 1]
    records = desired_records(result.catalog, result.assets)
    first = records["product-variations"]["v1"]["addons"]
    second = records["product-variations"]["v2"]["addons"]
    assert len(first) == 2 and len(second) == 1
    assert {r["$id"] for r in first}.isdisjoint(r["$id"] for r in second)
    assert records["architrave-types"][second[0]["$id"]]["additionalCost"] == "30"
    assert records["product-variations"]["v1"]["filter_values"]
    assert not records["product-variations"]["v2"]["filter_values"]


@pytest.mark.parametrize(
    "heading",
    [
        "Кромка",
        "Кромка (якщо немає вибору, залишити пустою)",
        "Кромка (або назва додаткового параметру)",
    ],
)
def test_edge_relation_is_separate_from_addons(tmp_path: Path, heading: str) -> None:
    from app.models.catalog import Addon
    from app.strapi.mapping import desired_records

    path = workbook(tmp_path, ["Тип покриття", "Покриття", "Скло", heading])
    result = parse_workbook(path)
    assert result.valid, result.issues
    variation = result.catalog.variations[0]
    variation.addons.append(Addon(name="Frame", group="Коробка", price_delta=Decimal(25)))
    records = desired_records(result.catalog, result.assets)
    mapped = records["product-variations"][variation.external_id]
    for field, collection, expected in [
        ("decor", "decors", "value0"),
        ("decor_color", "decor-colors", "value1"),
        ("decor_glass", "decor-glasses", "value2"),
    ]:
        assert records[collection][mapped[field]["$id"]]["name"] == expected
    edge = mapped["architrave_types"][0]
    addon = mapped["addons"][0]
    assert edge != addon
    assert records["architrave-types"][edge["$id"]]["name"] == "value3"
    assert records["architrave-types"][addon["$id"]]["name"] == "Frame"
    variation.attributes.pop(variation.attribute_definitions[-1].key)
    mapped = desired_records(result.catalog, result.assets)["product-variations"][
        variation.external_id
    ]
    assert mapped["architrave_types"] == []
    assert len(mapped["addons"]) == 1
