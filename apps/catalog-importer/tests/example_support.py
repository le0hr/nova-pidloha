"""Repair a COPY for tests only. Never change the personnel's source workbook."""

import base64
from pathlib import Path

from openpyxl import load_workbook


def corrected_example(directory: Path) -> Path:
    workbook = load_workbook(Path(__file__).parents[1] / "example.xlsx")
    doors = workbook["Двері"]
    for row in (8, 18, 22, 31, 37, 47, 51, 60):
        for column in range(1, 12):
            doors.cell(row, column).value = None
    for row in range(3, 61):
        if doors.cell(row, 1).value is not None:
            doors.cell(row, 11).value = "photo.png"
    floor = workbook["Підлога"]
    floor["D3"] = "Test floor"
    floor["C3"] = "Test brand"
    floor["I3"] = 2001
    floor["L3"] = 2001
    floor["L4"] = 2001
    floor["P2"] = "ID товару"
    floor["P3"] = 2001
    floor["G3"] = "photo.png"
    (directory / "photo.png").write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aY9sAAAAASUVORK5CYII="
        )
    )
    path = directory / "corrected-test.xlsx"
    workbook.save(path)
    workbook.close()
    return path
