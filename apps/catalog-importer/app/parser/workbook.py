"""Read workbooks without evaluating formulas or mutating the original file."""

import logging
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from app.config import Settings
from app.models.catalog import Addon, Attribute, Characteristic, Filter, Source
from app.parser.embedded_images import prepare_embedded, read_images
from app.parser.headers import Headers
from app.parser.products import parse_product
from app.parser.sections import discover
from app.parser.values import identifier, money, optional_text, scalar
from app.parser.variations import parse_variation
from app.validation.errors import Issue, ValidationResult
from app.validation.media import prepare_media
from app.validation.validator import Related, validate_catalog

logger = logging.getLogger(__name__)


def parse_workbook(path: Path, settings: Settings | None = None) -> ValidationResult:
    """Return all parse and semantic issues; never communicate with Strapi."""
    settings = settings or Settings()
    headers = Headers(settings)
    result = ValidationResult()
    related: list[Related] = []
    embedded = read_images(path, result)
    workbook = load_workbook(path, data_only=False)
    try:
        for sheet in workbook.worksheets:
            logger.info("parsing sheet %s", sheet.title)
            tables = discover(sheet, headers, result)
            if not tables and any(c.value is not None for row in sheet for c in row):
                result.error(
                    Source(sheet=sheet.title, row=1),
                    "header",
                    "No supported catalog table found. Add Product ID and Name, or a complete variation/side-section header.",
                )
            for table in tables:
                logger.info(
                    "detected %s table on %s row %s (%s attributes)",
                    table.kind,
                    sheet.title,
                    table.row,
                    len(table.attributes),
                )
                end = min(
                    (
                        t.row
                        for t in tables
                        if t.row > table.row and t.start <= table.end and t.end >= table.start
                    ),
                    default=sheet.max_row + 1,
                )
                combined_products: dict[str, Any] = {}
                empty_variation_rows: list[int] = []
                for row in range(table.row + 1, end):
                    cells = [sheet.cell(row, col) for col in range(table.start, table.end + 1)]
                    if all(c.value is None for c in cells):
                        if table.kind in {"variation", "combined"}:
                            empty_variation_rows.append(row)
                        continue
                    source = Source(
                        sheet=sheet.title, row=row, column=get_column_letter(table.start)
                    )
                    if any(c.data_type in {"f", "e"} for c in cells):
                        result.error(
                            source,
                            "cell",
                            "Formula or Excel error found. Replace it with a verified value before import.",
                        )
                        continue
                    # Only explicitly merged IDs are inherited. Ordinary blank cells are errors.
                    values: dict[str, Any] = {}
                    for key, col in table.fields.items():
                        value = sheet.cell(row, col).value
                        if value is None and key == "product_id":
                            for merged in sheet.merged_cells.ranges:
                                if (
                                    merged.min_col == merged.max_col == col
                                    and merged.min_row <= row <= merged.max_row
                                ):
                                    value = sheet.cell(merged.min_row, col).value
                        values[key] = value
                    try:
                        image_col = table.fields.get("image")
                        pictures = embedded.pop((sheet.title, row, image_col), []) if image_col else []
                        if pictures:
                            if len(pictures) != 1 or optional_text(values.get("image")):
                                raise ValueError(
                                    "Photo cell must contain one embedded image or one path, not both."
                                )
                            values["image"] = prepare_embedded(path, pictures[0], result)
                        if table.kind == "combined" and values.get("name") is None:
                            metadata = {"category", "brand", "description"}
                            ignored = [key for key in metadata if values.get(key) is not None]
                            ignored.extend(
                                display
                                for col, display in table.attributes.items()
                                if col < table.fields["variation_id"]
                                and sheet.cell(row, col).value is not None
                            )
                            if ignored:
                                result.error(
                                    source,
                                    "product",
                                    "Product metadata on a continuation row would be ambiguous. Move it to the named product row.",
                                    ignored,
                                )
                        if table.kind == "product" or (
                            table.kind == "combined" and values.get("name") is not None
                        ):
                            product_values = dict(values)
                            if table.kind == "combined":
                                product_values.pop("price", None)
                                product_values.pop("image", None)
                                product_values.pop("active", None)
                            product = parse_product(product_values, source, sheet.title)
                            for col, display in table.attributes.items():
                                if (
                                    "variation_id" in table.fields
                                    and col > table.fields["variation_id"]
                                ):
                                    continue
                                key = headers.attribute_key(display)
                                product.attribute_definitions.append(
                                    Attribute(display_name=display, key=key)
                                )
                                if sheet.cell(row, col).value is not None:
                                    product.attributes[key] = scalar(sheet.cell(row, col).value)
                            previous = combined_products.get(product.external_id)
                            if table.kind == "combined" and previous is not None:
                                if previous.model_dump(exclude={"source"}) != product.model_dump(
                                    exclude={"source"}
                                ):
                                    result.error(
                                        source,
                                        "product_id",
                                        f"Conflicting details for product {product.external_id}; compare with row {previous.source.row}.",
                                        product.external_id,
                                    )
                            else:
                                result.catalog.products.append(product)
                                combined_products[product.external_id] = product
                        if table.kind in {"variation", "combined"}:
                            variation = parse_variation(
                                values,
                                source.model_copy(
                                    update={
                                        "column": get_column_letter(table.fields["variation_id"])
                                    }
                                ),
                            )
                            for col, display in table.attributes.items():
                                if col < table.fields["variation_id"]:
                                    continue
                                key = headers.attribute_key(display)
                                variation.attribute_definitions.append(
                                    Attribute(display_name=display, key=key)
                                )
                                value = sheet.cell(row, col).value
                                if optional_text(value) is not None:
                                    variation.attributes[key] = scalar(value)
                            result.catalog.variations.append(variation)
                        elif table.kind == "filter":
                            result.catalog.filters.append(
                                Filter(
                                    target_external_id=identifier(
                                        values.get("target_id", values.get("product_id"))
                                    ),
                                    name=values["filter"],
                                    value=str(values["value"])
                                    if values.get("value") is not None
                                    else "",
                                    target_type=optional_text(values.get("target_type"))
                                    or table.target_type,
                                    source=source,
                                )
                            )
                        elif table.kind in {"characteristic", "addon", "description"}:
                            target = identifier(values.get("product_id"))
                            item: Characteristic | Addon | str
                            if table.kind == "characteristic":
                                item = Characteristic(
                                    name=values["characteristic"],
                                    value=str(values["value"])
                                    if values.get("value") is not None
                                    else "",
                                    unit=optional_text(values.get("unit")),
                                )
                            elif table.kind == "addon":
                                item = Addon(
                                    name=values["addon"],
                                    external_id=identifier(values["addon_id"])
                                    if values.get("addon_id") is not None
                                    else None,
                                    group=optional_text(values.get("group")),
                                    price=money(values["price"])
                                    if values.get("price") is not None
                                    else None,
                                    price_delta=money(values["price_delta"])
                                    if values.get("price_delta") is not None
                                    else None,
                                )
                            else:
                                item = (
                                    optional_text(
                                        values.get("description", values.get("short_description"))
                                    )
                                    or ""
                                )
                                if not item:
                                    raise ValueError("Enter description text.")
                            related.append(
                                Related(
                                    target,
                                    source,
                                    item,
                                    "short_description"
                                    if "short_description" in values
                                    else "description",
                                    target_type=table.target_type,
                                )
                            )
                    except ValidationError as exc:
                        for error in exc.errors(include_url=False):
                            field = str(error["loc"][0])
                            excel_field = {
                                "external_id": "variation_id"
                                if table.kind == "variation"
                                else "product_id",
                                "product_external_id": "product_id",
                            }.get(field, field)
                            at = source.model_copy(
                                update={
                                    "column": get_column_letter(
                                        table.fields.get(excel_field, table.start)
                                    )
                                }
                            )
                            result.error(at, field, error["msg"], error.get("input"))
                    except ValueError as exc:
                        result.error(source, "row", str(exc), values)
                if empty_variation_rows:
                    result.issues.append(
                        Issue(
                            severity="WARNING",
                            source=Source(sheet=sheet.title, row=empty_variation_rows[0]),
                            field="row",
                            message=(
                                f"{len(empty_variation_rows)} empty variation rows ignored in this "
                                "table; no records will be deleted."
                            ),
                        )
                    )
        for (title, row, col) in embedded:
            result.error(
                Source(sheet=title, row=row, column=get_column_letter(col)),
                "image",
                "Picture is not anchored to a Photo cell on a catalog data row. Move its top-left corner into that cell.",
            )
        validate_catalog(result, related)
        prepare_media(result, path, settings)
        return result
    finally:
        workbook.close()
