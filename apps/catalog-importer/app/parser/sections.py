"""Discover independent header bands, including horizontally adjacent tables."""

from dataclasses import dataclass

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models.catalog import Source
from app.parser.headers import Headers, normalize
from app.validation.errors import ValidationResult


@dataclass
class Table:
    kind: str
    row: int
    start: int
    end: int
    fields: dict[str, int]
    attributes: dict[int, str]
    target_type: str | None = None


def discover(sheet: Worksheet, headers: Headers, result: ValidationResult) -> list[Table]:
    tables: list[Table] = []
    for row_number, row in enumerate(sheet.iter_rows(), start=1):
        anchors = [
            c.column
            for c in row
            if headers.resolve(c.value) in {"product_id", "variation_id", "target_id"}
        ]
        if not anchors:
            continue
        # A data row may contain words resembling headers, so require another system label.
        if sum(headers.resolve(c.value) is not None for c in row) < 2:
            continue
        groups: list[tuple[int, int]] = []
        start = next((c.column for c in row if c.value is not None), 1)
        seen: set[str] = set()
        inside_variation = False
        for col, cell in enumerate(row, start=1):
            key = headers.resolve(cell.value)
            if col < start:
                continue
            boundary = (cell.value is None and not inside_variation) or (
                key in {"product_id", "variation_id", "target_id"} and key in seen
            )
            if boundary:
                if seen:
                    groups.append((start, col - 1))
                start = col + (cell.value is None)
                seen = set()
            if key:
                seen.add(key)
            if key == "variation_id":
                # Side sections have no variation attribute/price band.
                title = normalize(sheet.cell(row_number - 1, start).value) if row_number > 1 else ""
                inside_variation = not title.startswith(("додат", "фільтр"))
            if key == "price":
                inside_variation = False
        if seen:
            groups.append((start, sheet.max_column))
        for start, end in groups:
            while end > start and sheet.cell(row_number, end).value is None:
                end -= 1
            title = normalize(sheet.cell(row_number - 1, start).value) if row_number > 1 else ""
            context = next(
                (
                    kind
                    for prefix, kind in [
                        ("опис", "description"),
                        ("характеристик", "characteristic"),
                        ("додат", "addon"),
                        ("фільтр", "filter"),
                    ]
                    if title.startswith(prefix)
                ),
                None,
            )
            raw_names = {
                normalize(sheet.cell(row_number, col).value) for col in range(start, end + 1)
            }
            if "назва характеристики" in raw_names:
                context = "characteristic"
            if "збільшення ціни" in raw_names:
                context = "addon"
            # Filter tables have an ID, Name and Value, with no price boundary.
            if {"назва", "значення"} <= raw_names and any(
                headers.resolve(sheet.cell(row_number, col).value) == "variation_id"
                for col in range(start, end + 1)
            ):
                context = "filter"
            target_type = None
            fields: dict[str, int] = {}
            names: set[str] = set()
            r = row_number
            for col in range(start, end + 1):
                value = sheet.cell(r, col).value
                name, key = normalize(value), headers.resolve(value)
                if (
                    context == "characteristic"
                    and "назва характеристики" in raw_names
                    and name == "характеристика"
                ):
                    key = "value"
                if context == "addon":
                    if key == "variation_id":
                        key, target_type = "product_id", "variation"
                    key = {"category": "group", "name": "addon"}.get(key or "", key)
                if context == "filter":
                    if key == "variation_id":
                        key, target_type = "target_id", "variation"
                    elif key == "product_id":
                        key, target_type = "target_id", "product"
                    elif key == "name":
                        key = "filter"
                source = Source(sheet=sheet.title, row=r, column=get_column_letter(col))
                if name and (name in names or (key is not None and key in fields)):
                    result.error(
                        source, "header", "Duplicate header. Give each column a unique name.", value
                    )
                names.add(name)
                if key:
                    fields[key] = col
            if not fields.keys() & {"product_id", "variation_id", "target_id"}:
                continue
            kind = next(
                (
                    k
                    for k in ("characteristic", "addon", "filter", "description")
                    if k in fields and "variation_id" not in fields and "name" not in fields
                ),
                "product",
            )
            if "short_description" in fields and "name" not in fields:
                kind = "description"
            if context:
                kind = context
            elif "variation_id" in fields:
                kind = "combined" if "name" in fields else "variation"
            required = {
                "product": {"product_id", "name"},
                "combined": {"product_id", "variation_id", "name", "price"},
                "variation": {"product_id", "variation_id", "price"},
                "characteristic": {"product_id", "characteristic", "value"},
                "addon": {"product_id", "addon"},
                "description": {
                    "product_id",
                    "short_description" if "short_description" in fields else "description",
                },
                "filter": {"filter", "value"},
            }[kind]
            source = Source(sheet=sheet.title, row=r, column=get_column_letter(start))
            missing = required - fields.keys()
            if kind == "filter" and not fields.keys() & {"target_id", "product_id"}:
                missing.add("target_id")
            if missing:
                result.error(
                    source,
                    "header",
                    f"Missing headers: {', '.join(sorted(missing))}. Check this table's headings.",
                )
            allowed = {
                "product": {
                    "product_id",
                    "name",
                    "category",
                    "brand",
                    "description",
                    "short_description",
                    "price",
                    "image",
                    "active",
                },
                "combined": {
                    "product_id",
                    "name",
                    "category",
                    "brand",
                    "description",
                    "short_description",
                    "variation_id",
                    "price",
                    "image",
                    "sku",
                    "active",
                    "note",
                    "supplier_code",
                },
                "variation": {
                    "product_id",
                    "variation_id",
                    "price",
                    "image",
                    "sku",
                    "active",
                    "note",
                    "supplier_code",
                },
                "characteristic": {"product_id", "characteristic", "value", "unit"},
                "addon": {"product_id", "addon", "addon_id", "group", "price", "price_delta"},
                "description": {"product_id", "description", "short_description"},
                "filter": {"product_id", "target_id", "filter", "value", "target_type"},
            }[kind]
            for key in fields.keys() - allowed:
                result.error(
                    source,
                    key,
                    "System field is not supported in this section. Move it to the appropriate table.",
                )
            if "variation_id" in fields:
                for key in {
                    "product_id",
                    "name",
                    "category",
                    "brand",
                    "description",
                } & fields.keys():
                    if fields[key] > fields["variation_id"]:
                        result.error(source, key, "Place product fields before Variation ID.")
            attrs: dict[int, str] = {}
            left, right = fields.get("variation_id"), fields.get("price")
            if left is not None and right is not None:
                if right < left:
                    result.error(
                        source,
                        "price",
                        "Move Price after Variation ID; attributes belong between them.",
                    )
                for col in range(left + 1, right):
                    value = sheet.cell(r, col).value
                    at = Source(sheet=sheet.title, row=r, column=get_column_letter(col))
                    if not normalize(value):
                        result.error(
                            at, "header", "Enter an attribute name between Variation ID and Price."
                        )
                    elif headers.resolve(value):
                        result.error(
                            at,
                            "header",
                            "System fields cannot be inside the variation attribute region.",
                            value,
                        )
                    else:
                        attrs[col] = str(value).strip()
            for col in range(start, end + 1):
                value = sheet.cell(r, col).value
                if value is None or headers.resolve(value) or col in attrs:
                    continue
                at = Source(sheet=sheet.title, row=r, column=get_column_letter(col))
                suggestion = headers.suggestion(value)
                if suggestion:
                    result.error(
                        at,
                        "header",
                        f'Unknown reserved-looking header. Did you mean "{suggestion}"?',
                        value,
                    )
                elif kind in {"product", "combined"} and (left is None or col < left):
                    attrs[col] = str(value).strip()
                else:
                    result.error(
                        at,
                        "header",
                        "Unknown column outside the attribute region. Configure a system alias or move the attribute.",
                        value,
                    )
            keys = [headers.attribute_key(v) for v in attrs.values()]
            if len(set(keys)) != len(keys) or "" in keys:
                result.error(
                    source,
                    "header",
                    "Attribute names map to empty or duplicate keys. Rename them or configure ATTRIBUTE_KEYS.",
                )
            table = Table(kind, r, start, end, fields, attrs, target_type)
            tables.append(table)
            result.tables.append(
                f"{sheet.title}: {kind} {get_column_letter(start)}{r}:{get_column_letter(end)}{sheet.max_row}; {len(attrs)} attributes"
            )
    return tables
