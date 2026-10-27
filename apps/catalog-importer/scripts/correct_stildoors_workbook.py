"""Correct the supplied price workbook without reserializing its embedded media."""
from collections import defaultdict
from pathlib import Path
import re
from xml.etree import ElementTree as E
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter

ROOT = Path(__file__).resolve().parents[1]
source = ROOT / 'prices_with_stildoors_photos.xlsx'
target = ROOT / 'prices_with_stildoors_photos_import.xlsx'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
S = '{' + NS['s'] + '}'
E.register_namespace('', NS['s'])
w = load_workbook(source, read_only=True)
groups = defaultdict(list)
for row in w.worksheets[0].iter_rows(min_row=3, max_col=12):
    if row[0].value is not None:
        groups[str(row[0].value)].append(row)
previous = load_workbook(ROOT / 'prices_with_aliano_photos_import.xlsx', read_only=True)
old_filters = list(previous.worksheets[0].iter_rows(min_row=3, min_col=27, max_col=29, values_only=True))
new_filters = list(w.worksheets[0].iter_rows(min_row=3, min_col=26, max_col=28, values_only=True))
variation_products = {str(r[5]): str(r[0]) for r in w.worksheets[0].iter_rows(min_row=3, max_col=6, values_only=True) if r[5]}
assert len(old_filters) == len(new_filters)
for old, new in zip(old_filters, new_filters):
    assert old[1:] == new[1:]
    if old[0] is not None:
        assert variation_products[str(old[0])] == str(new[0])
previous.close()
changes = {}
for pid, rows in groups.items():
    if len({r[2].value for r in rows}) > 1:
        assert all(len({r[c].value for r in rows}) == 1 for c in (0, 1, 3, 4)), pid
        bases = {str(r[2].value).split(',')[0].strip() for r in rows}
        assert len(bases) == 1, (pid, bases)
        for row in rows:
            changes[row[0].row] = (next(iter(bases)), row[2].value)
w.close()


def shift_ref(ref):
    def sub(match):
        col = column_index_from_string(match[1])
        return get_column_letter(col + (col >= 10)) + match[2]
    return re.sub(r'([A-Z]+)([0-9]+)', sub, ref)


def set_text(cell, value):
    for child in list(cell):
        cell.remove(child)
    cell.set('t', 'inlineStr')
    E.SubElement(E.SubElement(cell, S + 'is'), S + 't').text = str(value)


with ZipFile(source) as zin, ZipFile(target, 'w') as zout:
    for info in zin.infolist():
        data = zin.read(info.filename)
        if info.filename == 'xl/worksheets/sheet1.xml':
            root = E.fromstring(data)
            for el in root.iter():
                for attr in ('r', 'ref', 'sqref', 'topLeftCell', 'activeCell'):
                    if attr in el.attrib:
                        el.set(attr, shift_ref(el.get(attr)))
            for col in root.findall('s:cols/s:col', NS):
                for key in ('min', 'max'):
                    n = int(col.get(key))
                    col.set(key, str(n + (n >= 10)))
            for row in root.findall('s:sheetData/s:row', NS):
                number = int(row.get('r'))
                if number == 2:
                    set_text(row.find('s:c[@r="AA2"]', NS), 'ID варіації')
                elif number >= 3 and number - 3 < len(old_filters) and old_filters[number - 3][0] is not None:
                    set_text(row.find(f's:c[@r="AA{number}"]', NS), old_filters[number - 3][0])
                if 'spans' in row.attrib:
                    row.attrib.pop('spans')
                if number == 2 or number in changes:
                    cell = E.Element(S + 'c', {'r': f'J{number}'})
                    if number == 2:
                        set_text(cell, 'Конструкція')
                    else:
                        base, original = changes[number]
                        set_text(cell, original)
                        set_text(row.find(f's:c[@r="C{number}"]', NS), base)
                    row.append(cell)
                    row[:] = sorted(row, key=lambda c: column_index_from_string(re.match(r'[A-Z]+', c.get('r'))[0]))
            data = E.tostring(root, encoding='utf-8', xml_declaration=True)
        elif info.filename == 'xl/drawings/drawing1.xml':
            root = E.fromstring(data)
            for col in root.iter('{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}col'):
                n = int(col.text)
                col.text = str(n + (n >= 9))
            data = E.tostring(root, encoding='utf-8', xml_declaration=True)
        elif info.filename == 'xl/worksheets/sheet2.xml':
            root = E.fromstring(data)
            # The existing CMS ID 2001 belongs to the hardware product Kofrad.
            cell = root.find('.//s:c[@r="A3"]', NS)
            assert cell.find('s:v', NS).text == '2001'
            set_text(cell, 'pidloha-2001')
            data = E.tostring(root, encoding='utf-8', xml_declaration=True)
        elif info.filename == 'xl/worksheets/sheet3.xml':
            root = E.fromstring(data)
            set_text(root.find('.//s:c[@r="P2"]', NS), 'ID товару')
            data = E.tostring(root, encoding='utf-8', xml_declaration=True)
        elif info.filename == 'xl/workbook.xml':
            root = E.fromstring(data)
            for sheet in root.findall('s:sheets/s:sheet', NS):
                if sheet.get('name') == 'фурнiтура':
                    sheet.set('name', 'Фурнітура')
            data = E.tostring(root, encoding='utf-8', xml_declaration=True)
        zout.writestr(info, data)
with ZipFile(source) as a, ZipFile(target) as b:
    for name in a.namelist():
        if name.startswith('xl/media/'):
            assert a.read(name) == b.read(name)
print(f'{target.name}: preserved media; {len(changes)} variation rows retain original construction; {len({str(groups[k][0][0].value) for k in groups if len({r[2].value for r in groups[k]}) > 1})} product groups corrected')
